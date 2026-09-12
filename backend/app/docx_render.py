"""決定書草稿的 Word 輸出:逐區塊印 build_decision_blocks,與 PDF 同一份版面定義。"""
import io

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt

from app.pdf_render import SPLIT, build_decision_blocks, NUMBERED_ITEM

_FONT = "新細明體"  # 字型不在閱讀端時由 Word 自行回退
_BODY_SIZE = Pt(11)
_HANGING = Pt(22)  # 條號「一、」的寬度,續行退到這裡


def _write(doc: Document, text: str, size: Pt):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = _FONT
    # 中文字走 eastAsia 那一欄,只設 font.name 的話 Word 只換英數字
    run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT)
    run.font.size = size
    return paragraph


def _usable_width(doc: Document):
    section = doc.sections[0]
    return section.page_width - section.left_margin - section.right_margin


def render_draft_docx(case) -> bytes:
    """依 build_decision_blocks(case) 產出 .docx bytes。與 render_draft_pdf 讀同一份版面——
    兩邊各自決定內容的下場是同一件案子下載到兩份不一樣的決定書。"""
    doc = Document()
    for kind, text in build_decision_blocks(case):
        if kind == "split":
            left, _, right = text.partition(SPLIT)
            paragraph = _write(doc, f"{left}\t{right}" if right else left, _BODY_SIZE)
            # 靠右交給定位點:用空白湊的位置換一台機器的字型就歪了
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                _usable_width(doc), WD_TAB_ALIGNMENT.RIGHT
            )
            continue
        for line in text.splitlines() or [""]:
            paragraph = _write(doc, line, _BODY_SIZE)
            if NUMBERED_ITEM.match(line):
                paragraph.paragraph_format.left_indent = _HANGING
                paragraph.paragraph_format.first_line_indent = -_HANGING
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
